"""Exportação da tabela "Compras do período" em Excel (.xlsx) e PDF."""

from datetime import datetime, timedelta, timezone
from io import BytesIO
from xml.sax.saxutils import escape

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from formatacao import brl, fmt_data

ROSA = colors.HexColor("#D6577A")
ROSA_CLARO = colors.HexColor("#FBE4E6")

CABECALHO = [
    "Compra", "Fornecedora", "Tipo de compra", "Data da compra",
    "Vencimento", "Situação", "Pago em", "Valor", "Desconto",
]


def _agora():
    return datetime.now(timezone(timedelta(hours=-3))).strftime("%d/%m/%Y %H:%M")


def gerar_excel_compras(linhas, periodo_txt, visao_txt, totais):
    """linhas: lista de dicts (id, fornecedora, tipo_compra, data_aceite,
    data_vencimento, situacao, data_pagamento, valor, desconto).
    totais: lista de (situacao, Decimal)."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Compras"

    ws.append([f"ENNE Brechó — Compras ({visao_txt})"])
    ws["A1"].font = Font(bold=True, size=14, color="D6577A")
    ws.append([f"Período: {periodo_txt}"])
    ws.append([f"Gerado em {_agora()}"])
    ws.append([])

    linha_cab = ws.max_row + 1
    ws.append(CABECALHO)
    for celula in ws[linha_cab]:
        celula.font = Font(bold=True, color="FFFFFF")
        celula.fill = PatternFill("solid", fgColor="D6577A")
        celula.alignment = Alignment(horizontal="center")

    for l in linhas:
        ws.append([
            l["id"], l["fornecedora"], l["tipo_compra"], l["data_aceite"],
            l["data_vencimento"], l["situacao"], l["data_pagamento"],
            float(round(l["valor"], 2)), float(round(l["desconto"], 2)),
        ])

    primeira = linha_cab + 1
    ultima = ws.max_row
    for r in range(primeira, ultima + 1):
        for col in (4, 5, 7):
            ws.cell(r, col).number_format = "DD/MM/YYYY"
        for col in (8, 9):
            ws.cell(r, col).number_format = '"R$" #,##0.00'

    ws.append([])
    for situacao, total in totais:
        ws.append([None, None, None, None, None, f"Total — {situacao}", None, float(round(total, 2))])
        r = ws.max_row
        ws.cell(r, 6).font = Font(bold=True)
        ws.cell(r, 8).font = Font(bold=True)
        ws.cell(r, 8).number_format = '"R$" #,##0.00'

    larguras = [9, 30, 28, 15, 13, 12, 13, 14, 12]
    for i, largura in enumerate(larguras, start=1):
        ws.column_dimensions[get_column_letter(i)].width = largura
    ws.freeze_panes = ws.cell(primeira, 1)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def gerar_pdf_compras(linhas, periodo_txt, visao_txt, totais):
    """Mesmos parâmetros do Excel. PDF em paisagem, com tabela paginada."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer, pagesize=landscape(A4),
        topMargin=1.5 * cm, bottomMargin=1.5 * cm, leftMargin=1.8 * cm, rightMargin=1.8 * cm,
    )
    estilos = getSampleStyleSheet()
    titulo = estilos["Heading1"]
    titulo.textColor = ROSA
    normal = estilos["Normal"]
    celula = ParagraphStyle("celula", parent=normal, fontSize=8, leading=10)
    celula_dir = ParagraphStyle("celula_dir", parent=celula, alignment=2)
    celula_cab = ParagraphStyle("celula_cab", parent=celula, textColor=colors.white, fontName="Helvetica-Bold")

    def p(texto, estilo=celula):
        return Paragraph(escape(str(texto)), estilo)

    elementos = [
        Paragraph(f"ENNE Brechó — Compras ({escape(visao_txt)})", titulo),
        Paragraph(f"<b>Período:</b> {escape(periodo_txt)} &nbsp;&nbsp; <b>Gerado em:</b> {_agora()}", normal),
        Spacer(1, 10),
    ]

    dados = [[p(c, celula_cab) for c in CABECALHO]]
    for l in linhas:
        dados.append([
            p(f"#{l['id']}"), p(l["fornecedora"]), p(l["tipo_compra"]), p(fmt_data(l["data_aceite"])),
            p(fmt_data(l["data_vencimento"])), p(l["situacao"]), p(fmt_data(l["data_pagamento"])),
            p(brl(l["valor"]), celula_dir),
            p(brl(l["desconto"]) if l["desconto"] > 0 else "—", celula_dir),
        ])

    larguras = [1.4, 6.0, 4.6, 2.4, 2.4, 2.3, 2.4, 2.4, 2.1]
    tabela = Table(dados, colWidths=[w * cm for w in larguras], repeatRows=1)
    tabela.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, 0), ROSA),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("GRID", (0, 0), (-1, -1), 0.5, colors.grey),
        ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, ROSA_CLARO]),
        ("LEFTPADDING", (0, 0), (-1, -1), 4), ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
    ]))
    elementos.append(tabela)
    elementos.append(Spacer(1, 10))
    for situacao, total in totais:
        elementos.append(Paragraph(f"<b>Total — {escape(situacao)}:</b> {brl(total)}", normal))

    doc.build(elementos)
    return buffer.getvalue()
